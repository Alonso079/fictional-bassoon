"""Añade la correspondencia funcional al Word existente sin tocar el original.

Ejecutar con el Python de las dependencias de documentos de Codex.
La fuente del anexo es docs/EQUIVALENCIAS_FUNCIONES_RUST.md.
"""

import argparse
import ast
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MAPPING = ROOT / "punishell-rust/docs/EQUIVALENCIAS_FUNCIONES_RUST.md"
SOURCE = ROOT / "INFORME_MAESTRO_PUNISHELL_RUST.docx"
OUTPUT = ROOT / "INFORME_MAESTRO_PUNISHELL_RUST_ACTUALIZADO.docx"


def check_coverage(text):
    tree = ast.parse((ROOT / "agente_python/agente_punishell.py").read_text(encoding="utf-8"))
    agent = next(node for node in tree.body if isinstance(node, ast.ClassDef))
    methods = [node.name for node in agent.body if isinstance(node, ast.FunctionDef)]
    missing = [name for name in methods if f"`{name}`" not in text]
    if missing:
        raise ValueError(f"Funciones Python sin correspondencia: {missing}")
    javascript = []
    for name in ["backend_php/includes/footer.php", "backend_php/panel_seguridad.php"]:
        javascript += re.findall(r"\bfunction\s+(\w+)\s*\(", (ROOT / name).read_text(encoding="utf-8"))
    missing_js = [name for name in javascript if f"`{name}`" not in text]
    if missing_js:
        raise ValueError(f"Funciones JavaScript sin correspondencia: {missing_js}")
    php = [path.name for path in (ROOT / "backend_php").glob("*.php")]
    missing_php = [name for name in php if name not in text]
    if missing_php:
        raise ValueError(f"Handlers PHP sin correspondencia: {missing_php}")
    print(f"Cobertura documental: {len(methods)} métodos Python, {len(javascript)} funciones JS y {len(php)} archivos PHP.")


def format_inline(paragraph, text, code_block=False):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        run = paragraph.add_run(part.strip("`") if part.startswith("`") else part.strip("*") if part.startswith("**") else part)
        if part.startswith("`") or code_block:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        if part.startswith("**"):
            run.bold = True


def table_style(table, headers, records):
    count = len(headers)
    if count == 5:
        widths = [0.48, 1.9, 2.7, 3.7, 0.8]
    elif count == 4:
        widths = [0.48, 2.45, 3.1, 3.55] if headers[0] == "ID" else [1.2, 2.5, 3.35, 2.55]
    else:
        widths = [2.25, 3.35, 4.0]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_pr.append(OxmlElement("w:tblHeader"))
    for index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        values = headers if index == 0 else records[index - 1]
        for column, cell in enumerate(row.cells):
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "163A63" if index == 0 else "F3F6FA" if index % 2 else "FFFFFF")
            tc_pr.append(shd)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), "85")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tc_pr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "bottom", "left", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "5")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            tc_pr.append(borders)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if headers[column] in ("ID", "Estado") else WD_ALIGN_PARAGRAPH.LEFT
            format_inline(p, values[column])
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if index == 0 else "26303A")
                if index == 0:
                    run.bold = True


def append_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    code = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code = not code
            i += 1
            continue
        if code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = i + 1 < len(lines) and not lines[i + 1].strip().startswith("```")
            format_inline(p, lines[i], code_block=True)
            i += 1
            continue
        if line.startswith("# "):
            doc.add_heading("Anexo C Correspondencia funcional y código Rust", level=1)
            i += 1
            continue
        if line.startswith("##"):
            level = min(len(line.split(" ", 1)[0]) - 1, 3)
            title = line.lstrip("#").strip()
            doc.add_heading(title, level=level + 1)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", cell) for cell in row):
                    rows.append(row)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table_style(table, rows[0], rows[1:])
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(4)
            spacer.paragraph_format.space_before = Pt(3)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            format_inline(p, line[2:])
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#", "|", "```", "- ")):
            paragraph_lines.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        format_inline(p, " ".join(paragraph_lines))


def update_report():
    text = MAPPING.read_text(encoding="utf-8")
    check_coverage(text)
    doc = Document(SOURCE)
    corrections = {
        "Puede capturar webcam, convertirla a JPEG/Base64 y enviarla al backend.":
        "Contiene capture_photo para webcam JPEG/Base64, pero no se llama desde gather_inventory ni run; no hay envío automático demostrado en ese flujo.",
        "La clave privada nunca debe estar en el agente.":
        "Las claves privadas del servidor y licenciamiento nunca deben estar en el agente; este sí mantiene su propia clave de identidad protegida localmente.",
        "La clave privada no se exporta ni se envía al servidor.":
        "La clave privada no se envía al servidor. DPAPI protege software; una clave TPM/CNG solo será no exportable si el proveedor y la implementación lo garantizan.",
        "Canonicalizar el payload antes de firmar y verificarlo; usar CBOR canónico o JSON canónico estable.":
        "Firmar y verificar los bytes exactos del sobre versionado del Anexo C; no reserializar el payload antes de comprobar la firma.",
        "Mensajes versionados, serialización canónica, errores":
        "Mensajes versionados, contrato de bytes firmados y errores",
        "Contrato canónico de comando + verificador Ed25519":
        "Contrato de bytes firmados + verificador Ed25519",
        "attestation": "challenge",
    }
    all_paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        all_paragraphs += [p for row in table.rows for cell in row.cells for p in cell.paragraphs]
    for p in all_paragraphs:
        for old, new in corrections.items():
            if old in p.text:
                p.text = p.text.replace(old, new)
    for p in doc.paragraphs:
        if p.text == "Resumen ejecutivo":
            note = p.insert_paragraph_before("Actualización al 11 de octubre de 2026. El Anexo C registra la equivalencia por función y el código Rust disponible. Se implementaron tres bibliotecas y 23 pruebas; el agente Windows y servidor aún requieren integración.")
            note.paragraph_format.space_after = Pt(10)
            break

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    append_markdown(doc, text)

    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        pr = doc.styles[name].element.find(qn("w:pPr"))
        if pr is not None:
            for border in list(pr.findall(qn("w:pBdr"))):
                pr.remove(border)
    for p in doc.paragraphs:
        if p.style.name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            pr = p._p.find(qn("w:pPr"))
            if pr is not None:
                for border in list(pr.findall(qn("w:pBdr"))):
                    pr.remove(border)
    doc.core_properties.title = "PuniShell y correspondencia funcional con Rust"
    doc.core_properties.subject = "Inventario de funciones y código de firmas licencias y permisos"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--check-coverage", action="store_true")
    args = parser.parse_args()
    if args.check_coverage:
        check_coverage(MAPPING.read_text(encoding="utf-8"))
    else:
        update_report()
